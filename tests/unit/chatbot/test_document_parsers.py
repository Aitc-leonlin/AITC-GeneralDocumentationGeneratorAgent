import io
import json
import unittest
from unittest.mock import patch

from docx import Document as DocxDocument
from langchain_core.documents import Document as LangChainDocument
from openpyxl import Workbook

from src.features.chatbot.services.document_parsers import (
    CsvParser,
    DocumentParseInput,
    DocumentParserError,
    DocumentParserService,
    DocxParser,
    ExcelParser,
    HTMLParser,
    PDFParser,
    TextParser,
    XBRLParser,
)


def parse_input(file_name: str, file_type: str, content: bytes) -> DocumentParseInput:
    return DocumentParseInput(
        file_name=file_name,
        file_type=file_type,
        content=content,
    )


class PDFParserTests(unittest.TestCase):
    @patch("src.features.chatbot.services.document_parsers.pdf_parser.PyPDFLoader")
    def test_preserves_pages_and_removes_repeated_headers_and_footers(self, loader_class):
        loader_class.return_value.load.return_value = [
            LangChainDocument(page_content="公司財務報告\n第一頁內容足以通過文字擷取門檻\nPage 1"),
            LangChainDocument(page_content="公司財務報告\n第二頁內容足以通過文字擷取門檻\nPage 2"),
        ]

        result = PDFParser().parse(parse_input("report.pdf", "pdf", b"pdf"))

        self.assertEqual(result.status, "parsed")
        self.assertEqual([block.page for block in result.blocks], [1, 2])
        self.assertEqual(result.blocks[0].content, "第一頁內容足以通過文字擷取門檻")
        self.assertEqual(result.blocks[1].content, "第二頁內容足以通過文字擷取門檻")

    @patch("src.features.chatbot.services.document_parsers.pdf_parser.PyPDFLoader")
    def test_returns_unsupported_scanned_pdf_for_nearly_empty_text(self, loader_class):
        loader_class.return_value.load.return_value = [
            LangChainDocument(page_content=" "),
            LangChainDocument(page_content="1"),
        ]

        result = PDFParser().parse(parse_input("scan.pdf", "pdf", b"pdf"))

        self.assertEqual(result.status, "unsupported_scanned_pdf")
        self.assertEqual(result.blocks, [])
        self.assertEqual(result.warnings[0].code, "unsupported_scanned_pdf")


class DocxParserTests(unittest.TestCase):
    def test_preserves_heading_paragraph_list_and_table(self):
        document = DocxDocument()
        document.add_heading("財務分析", level=1)
        document.add_paragraph("公司營收持續成長。")
        document.add_paragraph("注意短期借款", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "年度"
        table.rows[0].cells[1].text = "營收"
        table.rows[1].cells[0].text = "2025"
        table.rows[1].cells[1].text = "130"
        buffer = io.BytesIO()
        document.save(buffer)

        result = DocxParser().parse(parse_input("report.docx", "docx", buffer.getvalue()))

        self.assertEqual([block.type for block in result.blocks], [
            "heading",
            "paragraph",
            "list",
            "table",
        ])
        self.assertEqual(result.blocks[1].section, "財務分析")
        self.assertEqual(result.blocks[3].rows, [{"年度": "2025", "營收": "130"}])

    def test_uses_title_style_and_allows_document_without_heading(self):
        document = DocxDocument()
        document.add_paragraph("授信報告", style="Title")
        document.add_paragraph("沒有 Heading 也應保留此段落。")
        buffer = io.BytesIO()
        document.save(buffer)

        result = DocxParser().parse(parse_input("report.docx", "docx", buffer.getvalue()))

        self.assertTrue(result.blocks[0].metadata["isDocumentTitle"])
        self.assertEqual(result.blocks[1].type, "paragraph")


class HTMLParserTests(unittest.TestCase):
    def test_preserves_title_heading_hierarchy_and_table(self):
        content = b"""<!doctype html><html><head><title>Credit Report</title></head>
        <body><h1>Financial Analysis</h1><h2>Revenue</h2><p>Revenue grew.</p>
        <table><caption>Revenue Table</caption><tr><th>Year</th><th>Value</th></tr>
        <tr><td>2025</td><td>130</td></tr></table></body></html>"""

        result = HTMLParser().parse(parse_input("report.html", "html", content))

        self.assertTrue(result.blocks[0].metadata["isDocumentTitle"])
        paragraph = next(block for block in result.blocks if block.type == "paragraph")
        table = next(block for block in result.blocks if block.type == "table")
        self.assertEqual(paragraph.section, "Financial Analysis > Revenue")
        self.assertEqual(table.metadata["caption"], "Revenue Table")
        self.assertEqual(table.rows, [{"Year": "2025", "Value": "130"}])

    def test_preserves_elements_without_heading(self):
        result = HTMLParser().parse(
            parse_input("notes.html", "html", b"<article><p>Plain content</p></article>")
        )

        self.assertEqual(result.blocks[0].content, "Plain content")
        self.assertEqual(result.blocks[0].section, "")


class XBRLParserTests(unittest.TestCase):
    def test_preserves_complete_fact_semantics_and_report_metadata(self):
        content = b"""<?xml version="1.0" encoding="UTF-8"?>
        <xbrli:xbrl xmlns:xbrli="http://www.xbrl.org/2003/instance"
          xmlns:dei="urn:dei" xmlns:ifrs="urn:ifrs">
          <xbrli:context id="c1"><xbrli:entity><xbrli:identifier scheme="test">1234</xbrli:identifier></xbrli:entity>
          <xbrli:period><xbrli:instant>2025-12-31</xbrli:instant></xbrli:period></xbrli:context>
          <xbrli:unit id="TWD"><xbrli:measure>iso4217:TWD</xbrli:measure></xbrli:unit>
          <dei:EntityRegistrantName contextRef="c1">Example Corp</dei:EntityRegistrantName>
          <dei:DocumentType contextRef="c1">Annual Report</dei:DocumentType>
          <ifrs:Revenue contextRef="c1" unitRef="TWD" decimals="0">125000000</ifrs:Revenue>
        </xbrli:xbrl>"""

        result = XBRLParser().parse(parse_input("report.xbrl", "xbrl", content))
        revenue = next(row for row in result.blocks[0].rows if row["concept"] == "Revenue")

        self.assertEqual(revenue["period"], "2025-12-31")
        self.assertEqual(revenue["unit"], "iso4217:TWD")
        self.assertEqual(revenue["contextRef"], "c1")
        self.assertEqual(result.blocks[0].metadata["documentEntity"], "Example Corp")
        self.assertEqual(result.blocks[0].metadata["reportType"], "Annual Report")


class TextParserTests(unittest.TestCase):
    def test_preserves_markdown_section_hierarchy(self):
        content = """# 財務分析

整體表現穩定。

## 償債能力

流動比率改善。
"""

        result = TextParser().parse(parse_input("report.md", "md", content.encode()))

        paragraphs = [block for block in result.blocks if block.type == "paragraph"]
        self.assertEqual(paragraphs[0].section, "財務分析")
        self.assertEqual(paragraphs[1].section, "財務分析 > 償債能力")


class CsvParserTests(unittest.TestCase):
    def test_returns_row_based_structured_data(self):
        content = "年度,營收,淨利\n2025,130,3\n2024,120,2\n".encode()

        result = CsvParser().parse(parse_input("financial.csv", "csv", content))

        table = result.blocks[0]
        self.assertEqual(table.columns, ["年度", "營收", "淨利"])
        self.assertEqual(table.rows[0], {"年度": "2025", "營收": "130", "淨利": "3"})


class ExcelParserTests(unittest.TestCase):
    def test_preserves_sheet_columns_rows_and_irregular_warning(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "財務資料"
        worksheet.append(["財務資料表"])
        worksheet.append(["年度", "營收", "淨利"])
        worksheet.append([2025, 130, 3])
        worksheet.append([2024, 120])
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = ExcelParser().parse(parse_input("financial.xlsx", "xlsx", buffer.getvalue()))

        table = result.blocks[0]
        self.assertEqual(table.sheet, "財務資料")
        self.assertEqual(table.columns, ["年度", "營收", "淨利"])
        self.assertEqual(table.rows[0], {"年度": 2025, "營收": 130, "淨利": 3})
        self.assertIn("irregular_excel_layout", [warning.code for warning in result.warnings])


class DocumentParserServiceTests(unittest.TestCase):
    def test_returns_standard_text_document_elements(self):
        elements = DocumentParserService().parse(
            document_id="doc_001",
            file_name="report.md",
            content="# 財務狀況\n\n公司營運穩定。".encode(),
        )

        paragraph = next(
            element
            for element in elements
            if element.metadata["source_block_type"] == "paragraph"
        )
        self.assertEqual(paragraph.document_id, "doc_001")
        self.assertEqual(paragraph.file_name, "report.md")
        self.assertEqual(paragraph.file_type, "md")
        self.assertEqual(paragraph.content_type, "text")
        self.assertEqual(paragraph.content, "公司營運穩定。")
        self.assertEqual(paragraph.section, "財務狀況")
        self.assertIsNone(paragraph.page)
        self.assertIsNone(paragraph.sheet)

    def test_returns_standard_table_document_elements(self):
        elements = DocumentParserService().parse(
            document_id="doc_002",
            file_name="financial.csv",
            content="年度,營收\n2025,130\n".encode(),
        )

        self.assertEqual(len(elements), 1)
        table = elements[0]
        self.assertEqual(table.content_type, "table")
        self.assertEqual(
            json.loads(table.content),
            {
                "columns": ["年度", "營收"],
                "rows": [{"年度": "2025", "營收": "130"}],
            },
        )
        self.assertEqual(table.metadata["source_block_type"], "table")

    def test_rejects_unsupported_file_type(self):
        with self.assertRaises(DocumentParserError) as context:
            DocumentParserService().parse(
                document_id="doc_003",
                file_name="report.pptx",
                content=b"data",
            )

        self.assertEqual(context.exception.code, "unsupported_file_type")


if __name__ == "__main__":
    unittest.main()

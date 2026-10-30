"""將對話目前已有的規劃匯出成 DOCX，允許尚未完成大綱。"""

from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Cm, Pt

from src.features.report_generator.services.domain_report_service import _render_markdown_body

from .catalog import (
    load_content_type_templates,
    load_document_type_templates,
    load_section_type_templates,
)
from .docx_content_graph import build_docx_content
from .json_evidence import JsonEvidence
from .schemas import (
    GeneralDocumentOutline, GeneralDocumentPlan, PurposeDocumentOutline,
    PurposeDocumentPlan, ReportChatMessage,
)


def _add_readable_evidence(document: DocxDocument, evidence: JsonEvidence, evidence_ids: list[str]) -> None:
    """Show the cited JSON fields as source descriptions, never internal IDs."""
    labels = {
        "context": "案例情境",
        "facts": "案例事實",
        "sourceMaterials": "模擬來源描述",
        "assumptions": "案例假設",
    }
    selected = set(evidence_ids)
    document.add_paragraph("本章參考的虛構資料：")
    for item in evidence.items:
        if item.id in selected:
            label = labels[item.field.split("[", 1)[0]]
            document.add_paragraph(f"{label}：{item.text}", style="List Bullet")


def _add_generated_content(document: DocxDocument, paragraphs: list[str]) -> None:
    """Render AI Markdown as Word paragraphs and tables, including split table rows."""
    blocks: list[str] = []
    for paragraph in paragraphs:
        content = paragraph.strip()
        if not content:
            continue
        if blocks and content.startswith("|") and blocks[-1].rstrip().endswith("|"):
            blocks[-1] += "\n" + content
        else:
            blocks.append(content)
    if blocks:
        _render_markdown_body(document, "\n\n".join(blocks))


def export_current_draft(
    *,
    document_type_id: str | None,
    plan: GeneralDocumentPlan | None,
    outline: GeneralDocumentOutline | None,
    messages: list[ReportChatMessage],
    use_sample_data: bool = False,
    purpose_plan: PurposeDocumentPlan | None = None,
    purpose_outline: PurposeDocumentOutline | None = None,
) -> tuple[bytes, str]:
    """有大綱時依相關 JSON 或 LLM 撰寫；尚無大綱時匯出規劃進度。"""
    user_requests = [message.content for message in messages if message.role == "user"]
    if not user_requests:
        raise ValueError("請先描述文件需求，再下載目前版本")

    document_types = load_document_type_templates()
    section_types = load_section_type_templates()
    content_types = load_content_type_templates()
    purpose_references = (
        [item for purpose in purpose_plan.purposes for item in purpose.referenceTemplateIds]
        if purpose_plan else []
    )
    type_id = (
        purpose_references[0] if len(purpose_plan.purposes) == 1
        and len(purpose_references) == 1 else None
    ) if purpose_plan else (
        outline.documentTypeId if outline else plan.documentTypeId if plan else document_type_id
    )
    template = document_types.get(type_id) if type_id else None
    active_outline = purpose_outline or outline
    generated = (
        build_docx_content(active_outline, messages, document_type_id)
        if use_sample_data and active_outline else None
    )
    evidence = generated.get("evidence") if generated else None
    sample = evidence.sample if evidence else None
    answers = generated["answers"] if generated else None
    content_mode = generated["mode"] if generated else None
    title = (
        purpose_outline.title if purpose_outline else
        outline.title if outline else f"{template.name}（規劃草稿）" if template else "文件規劃草稿"
    )

    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin = page.bottom_margin = Cm(2)
    page.left_margin = page.right_margin = Cm(2.5)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal.font.size = Pt(11)

    document.add_heading(title, 0)
    version_text = (
        "版本：依目前大綱撰寫。"
        if content_mode == "outline_llm"
        else "版本：目前進度；缺少來源的內容標示為待補。"
    )
    document.add_paragraph(
        f"文件用途：{'、'.join(item.description for item in purpose_plan.purposes) if purpose_plan else type_id + ' ' + template.name if type_id and template else '待辨識'}\n"
        f"匯出日期：{datetime.now(ZoneInfo('Asia/Taipei')).date().isoformat()}\n"
        f"{version_text}"
    )

    document.add_heading("文件目的與範圍", level=1)
    if purpose_outline:
        document.add_paragraph(f"目的：{purpose_outline.objective}")
        document.add_paragraph(f"對象：{purpose_outline.subject}；期間：{purpose_outline.period}")
    elif outline:
        document.add_paragraph(f"目的：{outline.objective}")
        document.add_paragraph(f"對象：{outline.subject}；期間：{outline.period}")
    else:
        document.add_paragraph(f"預計目的：{template.purpose if template else '待確認'}")
        document.add_paragraph("對象與期間：待補")

    
    # if sample:
    #     print("==================sample:==================" + sample.model_dump_json())
    #     document.add_heading("預設測試資料", level=1)
    #     document.add_paragraph("以下資料為虛構測試案例，僅供示範文件產生。")
    #     document.add_paragraph(f"案例：{sample.title}；對象：{sample.subject}；期間：{sample.period}")
    #     document.add_paragraph(sample.context)
    #     table = document.add_table(rows=1, cols=2)
    #     table.style = "Light Shading Accent 1"
    #     table.rows[0].cells[0].text = "項目"
    #     table.rows[0].cells[1].text = "測試資料"
    #     for fact in sample.facts:
    #         cells = table.add_row().cells
    #         cells[0].text, cells[1].text = fact.label, fact.value

    document.add_heading("章節內容" if answers else "目前章節", level=1)
    if purpose_outline:
        for index, section in enumerate(purpose_outline.sections, start=1):
            document.add_heading(f"{index}. {section.heading}", level=2)
            if answers:
                _add_generated_content(document, answers[index - 1].paragraphs)
                if evidence and answers[index - 1].evidenceIds:
                    _add_readable_evidence(document, evidence, answers[index - 1].evidenceIds)
            else:
                document.add_paragraph(f"規劃重點：{section.focus}")
                document.add_paragraph("章節內文與來源資料：待補")
    elif purpose_plan:
        for index, section in enumerate(purpose_plan.sections, start=1):
            document.add_heading(f"{index}. {section.heading}", level=2)
            document.add_paragraph(f"規劃重點：{section.focusHint}")
            document.add_paragraph("章節內文與來源資料：待補")
    elif outline:
        for index, section in enumerate(outline.sections, start=1):
            document.add_heading(f"{index}. {section.heading}", level=2)
            if answers:
                _add_generated_content(document, answers[index - 1].paragraphs)
                if evidence and answers[index - 1].evidenceIds:
                    _add_readable_evidence(document, evidence, answers[index - 1].evidenceIds)
            else:
                document.add_paragraph(f"規劃重點：{section.focus}")
                document.add_paragraph("章節內文與來源資料：待補")
    elif plan:
        for index, section in enumerate(plan.sections, start=1):
            heading = section_types.get(section.sectionId)
            document.add_heading(
                f"{index}. {heading.name if heading else section.sectionId}", level=2
            )
            names = [
                content_types[item].name for item in section.contentTypeIds
                if item in content_types
            ]
            document.add_paragraph("預計內容：" + ("、".join(names) if names else "待規劃"))
            document.add_paragraph("章節重點、內文與來源資料：待補")
    elif template:
        for index, section_id in enumerate(template.requiredSectionIds, start=1):
            heading = section_types.get(section_id)
            document.add_heading(
                f"{index}. {heading.name if heading else section_id}", level=2
            )
            document.add_paragraph("章節內容：待規劃")
    else:
        document.add_paragraph("文件種類與章節：待規劃")

    document.add_heading("資料來源與限制", level=1)
    document.add_paragraph(
        "章節依目前大綱與使用者需求由 AI 撰寫，未使用預設測試 JSON；"
        "內容中的一般性說明不代表已完成個案查證。"
        if content_mode == "outline_llm" else
        purpose_outline.sourcePlan if purpose_outline else
        outline.sourcePlan if outline else "尚未確認資料來源。"
    )
    if sample:
        document.add_paragraph(sample.sourceNote)
        document.add_paragraph("案例情境：" + sample.context, style="List Bullet")
        for source in sample.sourceMaterials:
            document.add_paragraph(f"模擬來源描述：{source.title}：{source.content}", style="List Bullet")
        for assumption in sample.assumptions:
            document.add_paragraph(f"案例假設：{assumption}", style="List Bullet")
    document.add_paragraph(
        "本文件依目前對話與已完成的大綱規劃匯出；"
        + ("章節中的案例事實來自虛構測試 JSON；" if sample else "")
        + "未經資料查證的內容不作為事實或結論。"
    )
    document.add_heading("使用者於此聊天視窗提過的報告需求", level=1)
    for request in user_requests:
        document.add_paragraph(request, style="List Bullet")

    stream = BytesIO()
    document.save(stream)
    safe_type_id = type_id if type_id in document_types else "general"
    return stream.getvalue(), f"{safe_type_id}_current_draft.docx"

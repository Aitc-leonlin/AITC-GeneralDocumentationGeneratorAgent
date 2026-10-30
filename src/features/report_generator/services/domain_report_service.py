"""既存固定格式報告的資料結構與 DOCX 輸出。"""

import json
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from random import SystemRandom
from typing import Literal
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt
from markdown_it.token import Token
from pydantic import BaseModel, Field

from src.features.report_generator.services.sample_data_factory import randomize_report_sample


SAMPLES_PATH = Path(__file__).resolve().parents[1] / "data" / "domain_report_samples.json"
DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
Domain = Literal["medical", "finance", "education", "food"]
Method = Literal["ai_markdown", "standard_docx"]
DOMAIN_LABELS = {
    "medical": "醫療",
    "finance": "金融",
    "education": "教育",
    "food": "美食",
}


class ReportMetric(BaseModel):
    label: str = Field(min_length=1, max_length=100)
    value: str = Field(min_length=1, max_length=200)


class ReportSection(BaseModel):
    heading: str = Field(min_length=1, max_length=100)
    details: str = Field(min_length=1, max_length=4000)
    bullets: list[str] = Field(default_factory=list, max_length=12)


class ReportSourceBlock(BaseModel):
    title: str = Field(min_length=1, max_length=100)
    content: str = Field(min_length=1)


class ReportLLMTrace(BaseModel):
    stage: str = Field(min_length=1, max_length=100)
    requestJson: str = Field(min_length=1)
    responseJson: str = Field(min_length=1)


class DomainReportData(BaseModel):
    domain: Domain
    title: str = Field(min_length=1, max_length=150)
    subject: str = Field(min_length=1, max_length=150)
    period: str = Field(min_length=1, max_length=100)
    author: str = Field(default="", max_length=100)
    summary: str = Field(min_length=1, max_length=4000)
    metrics: list[ReportMetric] = Field(min_length=1, max_length=30)
    sections: list[ReportSection] = Field(min_length=1, max_length=12)
    sourceNote: str = Field(min_length=1, max_length=2000)
    sourceBlocks: list[ReportSourceBlock] = Field(default_factory=list)


class GenerateDomainReportRequest(BaseModel):
    method: Method
    data: DomainReportData


def load_sample_reports() -> dict[str, dict]:
    with SAMPLES_PATH.open(encoding="utf-8") as sample_file:
        samples = json.load(sample_file)
    randomizer = SystemRandom()
    return {
        domain: DomainReportData.model_validate(
            randomize_report_sample(domain, data, randomizer)
        ).model_dump()
        for domain, data in samples.items()
    }


def load_default_reports() -> dict[str, DomainReportData]:
    """讀取未經隨機化的固定 JSON，供快捷提問引用。"""
    with SAMPLES_PATH.open(encoding="utf-8") as sample_file:
        reports = json.load(sample_file)
    return {
        domain: DomainReportData.model_validate(data)
        for domain, data in reports.items()
    }


def llm_trace_blocks(traces: list[ReportLLMTrace], start: int = 1) -> list[ReportSourceBlock]:
    blocks = []
    stage_labels = {
        "report_outline.classify_question": "大綱問題分類",
        "report_outline.create": "大綱產生與修訂",
        "report_outline.extract_user_metrics": "使用者數值擷取",
        "report_outline.analyze": "報告內容分析",
    }
    for index, trace in enumerate(traces, start=start):
        stage = trace.stage
        for key, label in stage_labels.items():
            stage = stage.replace(key, label)
        stage = stage[:60]
        try:
            request_payload = json.loads(trace.requestJson)
            prompt_text = "\n\n".join(
                f"{message.get('role', 'user').upper()}：\n{message.get('content', '')}"
                for message in request_payload.get("messages", [])
            ) or "（無 Prompt 文字）"
        except (ValueError, AttributeError, TypeError):
            prompt_text = "（無法解析 REQUEST JSON）"
        blocks.extend([
            ReportSourceBlock(
                title=f"LLM 呼叫 {index}｜{stage}｜REQUEST JSON",
                content=trace.requestJson,
            ),
            ReportSourceBlock(
                title=f"LLM 呼叫 {index}｜{stage}｜PROMPT 原文",
                content=prompt_text,
            ),
            ReportSourceBlock(
                title=f"LLM 呼叫 {index}｜{stage}｜RESPONSE JSON",
                content=trace.responseJson,
            ),
        ])
    return blocks


def _raw_llm_response(message) -> dict:
    """保留 LangChain 回傳的原始訊息欄位；不包含 API 金鑰或 HTTP 標頭。"""
    if hasattr(message, "model_dump"):
        try:
            return message.model_dump(mode="json", exclude_none=True)
        except Exception:
            return {
                "content": getattr(message, "content", ""),
                "additional_kwargs": getattr(message, "additional_kwargs", {}),
                "response_metadata": getattr(message, "response_metadata", {}),
                "tool_calls": getattr(message, "tool_calls", []),
            }
    return {"content": str(message)}


def find_missing_metric_claims(text: str, metrics: list[ReportMetric]) -> list[str]:
    """找出與已知指標值直接矛盾的「缺少數值」敘述。"""
    cleaned = re.sub(r"[*_`]", "", text)
    missing = r"(?:沒有|缺少|未提供|查無|無法取得|缺乏)"
    labels = []
    for metric in metrics:
        if metric.value == "待提供":
            continue
        label = re.escape(metric.label)
        patterns = (
            rf"{label}[\s：:]*的?[\s：:]*{missing}[\s：:]*(?:相關|詳細|完整|可用)?[\s：:]*(?:資料|數據|數值)",
            rf"{missing}[\s：:]*{label}[\s：:]*的?[\s：:]*(?:相關|詳細|完整|可用)?[\s：:]*(?:資料|數據|數值)",
            rf"{label}[\s：:]*的?[\s：:]*(?:資料|數據|數值)[\s：:]*(?:不足|缺漏|未提供)",
        )
        if any(re.search(pattern, cleaned) for pattern in patterns):
            labels.append(metric.label)
    return labels


def _set_font(run, *, size: int = 11, bold: bool = False, color: str = "253341") -> None:
    run.font.name = "Microsoft JhengHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")


def _paragraph(
    document, text: str, *, style: str | None = None, size: int = 11,
    bold: bool = False, color: str = "253341",
):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(7)
    _set_font(paragraph.add_run(text), size=size, bold=bold, color=color)
    return paragraph


def _heading(document, text: str, level: int = 1):
    paragraph = _paragraph(
        document, text, style=f"Heading {level}", size=15 if level == 1 else 12,
        bold=True, color="176B69",
    )
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _new_document(data: DomainReportData, method_label: str):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Microsoft JhengHei"
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.line_spacing = 1.35
    normal_style.paragraph_format.space_after = Pt(7)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(header.add_run(f"{DOMAIN_LABELS[data.domain]}領域｜資料報告"), size=9, color="64748B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(footer.add_run("報告文件  •  "), size=9, color="64748B")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    title = _paragraph(document, data.title, size=22, bold=True, color="143F40")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = _paragraph(
        document, f"{DOMAIN_LABELS[data.domain]}｜{method_label}", size=10, color="64748B"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    metadata = document.add_table(rows=0, cols=2)
    metadata.style = "Light Shading Accent 1"
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in (
        ("對象／主題", data.subject),
        ("報告期間", data.period),
        ("建立者", data.author or "未填寫"),
        ("建立日期", datetime.now(ZoneInfo("Asia/Taipei")).date().isoformat()),
    ):
        cells = metadata.add_row().cells
        _set_font(cells[0].paragraphs[0].add_run(label), bold=True)
        _set_font(cells[1].paragraphs[0].add_run(value))
    document.add_paragraph()
    return document


def _render_standard(data: DomainReportData) -> bytes:
    document = _new_document(data, "標準 DOCX 版")
    _heading(document, "一、摘要")
    _paragraph(document, data.summary)

    _heading(document, "二、關鍵指標")
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ("指標", "數值")):
        _set_font(cell.paragraphs[0].add_run(text), bold=True)
    for metric in data.metrics:
        cells = table.add_row().cells
        _set_font(cells[0].paragraphs[0].add_run(metric.label))
        _set_font(cells[1].paragraphs[0].add_run(metric.value))

    _heading(document, "三、內容分析")
    for index, section in enumerate(data.sections, start=1):
        _heading(document, f"3.{index} {section.heading}", level=2)
        _paragraph(document, section.details)
        for bullet in section.bullets:
            _paragraph(document, bullet, style="List Bullet")

    _heading(document, "四、資料來源與限制")
    _paragraph(document, data.sourceNote)
    for index, block in enumerate(data.sourceBlocks, start=1):
        _heading(document, f"4.{index} {block.title}", level=2)
        is_json = block.title.endswith("JSON")
        paragraph = _paragraph(document, block.content, size=8 if is_json else 10)
        if is_json:
            for run in paragraph.runs:
                run.font.name = "Consolas"
        paragraph.paragraph_format.left_indent = Cm(0.35)
        paragraph.paragraph_format.right_indent = Cm(0.35)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(12)
        _shade_paragraph(paragraph, "F1F7F6")
    return _to_bytes(document)


def _generate_markdown(
    data: DomainReportData,
    traces: list[ReportLLMTrace] | None = None,
) -> str:
    # AI 僅能使用本次提交的 JSON；Markdown 會再轉成 DOCX 原生樣式。
    from langchain_core.messages import HumanMessage, SystemMessage
    from src.features.chatbot.core.providers.chat_openAI_provider import chat_model, get_message_text

    prompt = (
        "根據下方 JSON 寫一份繁體中文報告。只使用 JSON 所含的事實與數字，不補造資料。"
        "只輸出標準 Markdown，不要 HTML、包住全文的程式碼圍欄或前言。"
        "請依序使用且只使用四個二級標題："
        "## 一、摘要、## 二、關鍵指標、## 三、內容分析、## 四、資料來源與限制。"
        "第三節每個 sections 項目用三級標題，metrics 請用 Markdown 表格；"
        "第三章須忠實保留 sections 中引用的指標名稱、數值與分析，"
        "已有數值的指標不可改寫成缺少資料。"
        "第四章的資料來源區塊會由程式填入，請只輸出第四章標題，勿自行改寫來源資料。\n\n"
        f"JSON:\n{data.model_dump_json(indent=2, exclude={'sourceBlocks'})}"
    )
    system_prompt = "你是文件編輯。嚴格依使用者提供的 JSON 整理報告，禁止自行推測。"
    request_json = json.dumps({
        "captureLevel": "LangChain invoke input; provider HTTP body may add parameters",
        "operation": "report_markdown.format",
        "model": getattr(chat_model, "model_name", None),
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": prompt},
        ],
    }, ensure_ascii=False, indent=2)
    response = chat_model.invoke([
        SystemMessage(content=system_prompt),
        HumanMessage(content=prompt),
    ])
    if traces is not None:
        traces.append(ReportLLMTrace(
            stage="Markdown 排版",
            requestJson=request_json,
            responseJson=json.dumps({
                "captureLevel": "LangChain raw AIMessage",
                "raw": _raw_llm_response(response),
            }, ensure_ascii=False, indent=2, default=str),
        ))
    markdown = get_message_text(response).strip()
    if not markdown:
        raise ValueError("AI 未回傳 Markdown 內容")
    lines = markdown.splitlines()
    if len(lines) >= 2 and lines[0].strip().lower() in {"```markdown", "```md", "```"} and lines[-1].strip() == "```":
        markdown = "\n".join(lines[1:-1]).strip()
    return markdown


def _render_inline(paragraph, children: list[Token], *, size: int = 11, color: str = "253341", bold: bool = False) -> None:
    strong_depth = 0
    emphasis_depth = 0
    strike_depth = 0
    link_depth = 0
    for child in children:
        token_type = child.type
        if token_type == "strong_open":
            strong_depth += 1
            continue
        if token_type == "strong_close":
            strong_depth = max(0, strong_depth - 1)
            continue
        if token_type == "em_open":
            emphasis_depth += 1
            continue
        if token_type == "em_close":
            emphasis_depth = max(0, emphasis_depth - 1)
            continue
        if token_type == "s_open":
            strike_depth += 1
            continue
        if token_type == "s_close":
            strike_depth = max(0, strike_depth - 1)
            continue
        if token_type == "link_open":
            link_depth += 1
            continue
        if token_type == "link_close":
            link_depth = max(0, link_depth - 1)
            continue
        if token_type == "hardbreak":
            paragraph.add_run().add_break()
            continue
        if token_type == "softbreak":
            paragraph.add_run(" ")
            continue
        if token_type not in {"text", "code_inline", "image"}:
            continue

        content = f"〔圖片：{child.content}〕" if token_type == "image" else child.content
        run = paragraph.add_run(content)
        _set_font(
            run, size=size, bold=bold or strong_depth > 0,
            color="176B69" if link_depth else color,
        )
        run.italic = emphasis_depth > 0
        run.font.strike = strike_depth > 0
        run.underline = link_depth > 0
        if token_type == "code_inline":
            run.font.name = "Consolas"
            run.font.size = Pt(max(9, size - 1))


def _shade_cell(cell, color: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), color)


def _shade_paragraph(paragraph, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    paragraph._p.get_or_add_pPr().append(shading)


def _render_markdown_table(document, tokens: list[Token], start: int) -> int:
    rows: list[tuple[bool, list[list[Token]]]] = []
    cells: list[list[Token]] = []
    header = False
    current_cell: list[Token] = []
    index = start + 1
    while index < len(tokens) and tokens[index].type != "table_close":
        token = tokens[index]
        if token.type == "tr_open":
            cells = []
            header = False
        elif token.type in {"th_open", "td_open"}:
            current_cell = []
            header = header or token.type == "th_open"
        elif token.type == "inline":
            current_cell = token.children or []
        elif token.type in {"th_close", "td_close"}:
            cells.append(current_cell)
        elif token.type == "tr_close":
            rows.append((header, cells))
        index += 1

    if rows:
        column_count = max(len(cells) for _, cells in rows)
        table = document.add_table(rows=0, cols=column_count)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for row_index, (is_header, row_cells) in enumerate(rows):
            row = table.add_row()
            row_properties = row._tr.get_or_add_trPr()
            row_properties.append(OxmlElement("w:cantSplit"))
            if is_header:
                repeat_header = OxmlElement("w:tblHeader")
                repeat_header.set(qn("w:val"), "true")
                row_properties.append(repeat_header)
            for column_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _shade_cell(cell, "176B69" if is_header else ("F1F7F6" if row_index % 2 == 0 else "FFFFFF"))
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.keep_together = True
                if column_index > 0 and not is_header:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if column_index < len(row_cells):
                    _render_inline(
                        paragraph, row_cells[column_index], size=10,
                        color="FFFFFF" if is_header else "253341", bold=is_header,
                    )
        document.add_paragraph().paragraph_format.space_after = Pt(2)
    return index + 1


def _render_markdown_body(document, markdown: str) -> None:
    """將 AI Markdown 的結構轉為 Word 標題、段落、表格與清單。"""
    tokens = MarkdownIt("default").enable("table").parse(markdown)
    list_stack: list[str] = []
    quote_depth = 0
    index = 0
    while index < len(tokens):
        token = tokens[index]
        token_type = token.type
        if token_type == "table_open":
            index = _render_markdown_table(document, tokens, index)
            continue
        if token_type in {"bullet_list_open", "ordered_list_open"}:
            list_stack.append(token_type)
        elif token_type in {"bullet_list_close", "ordered_list_close"}:
            if list_stack:
                list_stack.pop()
        elif token_type == "blockquote_open":
            quote_depth += 1
        elif token_type == "blockquote_close":
            quote_depth = max(0, quote_depth - 1)
        elif token_type in {"heading_open", "paragraph_open"}:
            inline = tokens[index + 1] if index + 1 < len(tokens) else None
            if token_type == "heading_open":
                markdown_level = int(token.tag[1:])
                heading_level = min(3, max(1, markdown_level - 1))
                paragraph = document.add_paragraph(style=f"Heading {heading_level}")
                paragraph.paragraph_format.space_before = Pt(15 if heading_level == 1 else 10)
                paragraph.paragraph_format.space_after = Pt(7)
                paragraph.paragraph_format.keep_with_next = True
                _render_inline(
                    paragraph, inline.children or [] if inline else [],
                    size=15 if heading_level == 1 else 12,
                    color="176B69", bold=True,
                )
            else:
                style = "List Bullet" if list_stack and list_stack[-1] == "bullet_list_open" else (
                    "List Number" if list_stack else ("Quote" if quote_depth else None)
                )
                paragraph = document.add_paragraph(style=style)
                paragraph.paragraph_format.space_after = Pt(6)
                if list_stack:
                    paragraph.paragraph_format.left_indent = Cm(0.55 * len(list_stack))
                elif quote_depth:
                    paragraph.paragraph_format.left_indent = Cm(0.7)
                    paragraph.paragraph_format.right_indent = Cm(0.4)
                    _shade_paragraph(paragraph, "F1F7F6")
                _render_inline(paragraph, inline.children or [] if inline else [])
            index += 3
            continue
        elif token_type in {"fence", "code_block"}:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(10)
            _shade_paragraph(paragraph, "F1F5F9")
            for line_index, line in enumerate(token.content.rstrip("\n").split("\n")):
                if line_index:
                    paragraph.add_run().add_break()
                run = paragraph.add_run(line)
                _set_font(run, size=9, color="334155")
                run.font.name = "Consolas"
        elif token_type == "hr":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:color"), "A9C8C3")
            bottom.set(qn("w:sz"), "6")
            border.append(bottom)
            paragraph._p.get_or_add_pPr().append(border)
        index += 1


def _render_ai_markdown(
    data: DomainReportData,
    traces: list[ReportLLMTrace] | None = None,
) -> bytes:
    previous_trace_count = len(traces) if traces is not None else 0
    markdown = _generate_markdown(data, traces)
    if traces is not None:
        data.sourceBlocks.extend(llm_trace_blocks(
            traces[previous_trace_count:], start=previous_trace_count + 1
        ))
    required_headings = ["一、摘要", "二、關鍵指標", "三、內容分析", "四、資料來源與限制"]
    actual_headings = [
        line[3:].strip()
        for line in markdown.splitlines()
        if line.startswith("## ")
    ]
    if actual_headings != required_headings:
        # AI 偶爾會改動章節；仍以既有標準模板輸出完整四章節。
        return _render_standard(data)
    analysis_body = markdown.split("## 三、內容分析", 1)[1].split("## 四、資料來源與限制", 1)[0]
    original_analysis = "\n".join(
        section.details + "\n" + "\n".join(section.bullets)
        for section in data.sections
    )
    normalized_body = "".join(analysis_body.split())
    for metric in data.metrics:
        if metric.value == "待提供":
            continue
        if metric.label in original_analysis and metric.value in original_analysis:
            if "".join(metric.label.split()) not in normalized_body or "".join(metric.value.split()) not in normalized_body:
                return _render_standard(data)
    if find_missing_metric_claims(analysis_body, data.metrics):
        return _render_standard(data)
    source_chapter = ["## 四、資料來源與限制", "", data.sourceNote]
    for index, block in enumerate(data.sourceBlocks, start=1):
        source_chapter.extend([
            "",
            f"### 4.{index} {block.title}",
            "",
            *[f"> {line}" for line in block.content.splitlines()],
        ])
    markdown = markdown.split("## 四、資料來源與限制", 1)[0].rstrip() + "\n\n" + "\n".join(source_chapter)
    document = _new_document(data, "AI Markdown 排版版")
    _render_markdown_body(document, markdown)
    return _to_bytes(document)


def _to_bytes(document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def generate_domain_report(
    request: GenerateDomainReportRequest,
    traces: list[ReportLLMTrace] | None = None,
) -> tuple[bytes, str]:
    data = request.data
    output = _render_ai_markdown(data, traces) if request.method == "ai_markdown" else _render_standard(data)
    filename = f"{data.domain}_{request.method}_report.docx"
    return output, filename

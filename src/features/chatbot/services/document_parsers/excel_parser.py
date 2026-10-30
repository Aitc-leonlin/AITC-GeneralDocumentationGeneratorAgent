from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from src.features.chatbot.services.document_parsers.base import (
    BaseDocumentParser,
    DocumentParseInput,
    DocumentParserError,
)
from src.features.chatbot.services.document_parsers.common import (
    normalize_columns,
    serializable_cell_value,
)
from src.features.chatbot.services.document_parsers.models import (
    ParsedContentBlock,
    ParsedDocument,
    ParserWarning,
)


HEADER_SCAN_ROW_LIMIT = 20


class ExcelParser(BaseDocumentParser):
    parser_name = "excel"
    supported_extensions = frozenset({"xlsx"})

    def parse(self, source: DocumentParseInput) -> ParsedDocument:
        if not source.content:
            raise DocumentParserError("empty_document", "XLSX 文件沒有內容。")
        try:
            workbook = load_workbook(
                io.BytesIO(source.content),
                read_only=True,
                data_only=False,
            )
        except Exception as exc:
            raise DocumentParserError("invalid_xlsx", f"XLSX 解析失敗：{exc}") from exc

        blocks: list[ParsedContentBlock] = []
        warnings: list[ParserWarning] = []
        try:
            for worksheet in workbook.worksheets:
                sheet_block, sheet_warnings = self._parse_sheet(
                    worksheet.title,
                    [
                        [serializable_cell_value(cell) for cell in row]
                        for row in worksheet.iter_rows(values_only=True)
                    ],
                )
                if sheet_block is not None:
                    blocks.append(sheet_block)
                warnings.extend(sheet_warnings)
        finally:
            workbook.close()

        if not blocks:
            warnings.append(
                ParserWarning(code="empty_document", message="XLSX 沒有可解析的資料。")
            )

        return ParsedDocument(
            fileName=source.file_name,
            fileType=source.extension,
            parser=self.parser_name,
            blocks=blocks,
            warnings=warnings,
        )

    @classmethod
    def _parse_sheet(
        cls,
        sheet_name: str,
        matrix: list[list[Any]],
    ) -> tuple[ParsedContentBlock | None, list[ParserWarning]]:
        warnings: list[ParserWarning] = []
        non_empty_rows = [
            (index, row)
            for index, row in enumerate(matrix, start=1)
            if any(value not in (None, "") for value in row)
        ]
        if not non_empty_rows:
            return None, [
                ParserWarning(
                    code="empty_excel_sheet",
                    message="工作表沒有資料。",
                    context={"sheet": sheet_name},
                )
            ]

        header_position = cls._detect_header_position(non_empty_rows)
        header_row_number, header_values = non_empty_rows[header_position]
        last_header_column = max(
            index for index, value in enumerate(header_values) if value not in (None, "")
        )
        header_values = header_values[: last_header_column + 1]
        columns, columns_changed = normalize_columns(header_values)

        if header_position > 0:
            warnings.append(
                ParserWarning(
                    code="excel_leading_rows_skipped",
                    message="工作表標題列之前包含額外內容，已從推定的欄位列開始解析。",
                    context={"sheet": sheet_name, "headerRow": header_row_number},
                )
            )
        if columns_changed:
            warnings.append(
                ParserWarning(
                    code="normalized_excel_columns",
                    message="工作表包含空白或重複欄名，已產生安全欄名。",
                    context={"sheet": sheet_name, "headerRow": header_row_number},
                )
            )

        rows: list[dict[str, Any]] = []
        irregular_rows: list[int] = []
        for row_number, values in non_empty_rows[header_position + 1 :]:
            meaningful_width = max(
                (index + 1 for index, value in enumerate(values) if value not in (None, "")),
                default=0,
            )
            if meaningful_width != len(columns):
                irregular_rows.append(row_number)
            padded = [*values[: len(columns)], *([None] * max(0, len(columns) - len(values)))]
            rows.append(dict(zip(columns, padded)))

        if irregular_rows:
            warnings.append(
                ParserWarning(
                    code="irregular_excel_layout",
                    message="工作表部分資料列的有效欄位範圍與標題列不同，請確認解析結果。",
                    context={
                        "sheet": sheet_name,
                        "rows": irregular_rows[:100],
                        "additionalRows": max(0, len(irregular_rows) - 100),
                    },
                )
            )

        return (
            ParsedContentBlock(
                type="table",
                sheet=sheet_name,
                columns=columns,
                rows=rows,
                metadata={"headerRow": header_row_number, "rowCount": len(rows)},
            ),
            warnings,
        )

    @staticmethod
    def _detect_header_position(non_empty_rows: list[tuple[int, list[Any]]]) -> int:
        candidates = non_empty_rows[:HEADER_SCAN_ROW_LIMIT]
        for position, (_, row) in enumerate(candidates):
            populated_columns = {
                index for index, value in enumerate(row) if value not in (None, "")
            }
            if len(populated_columns) < 2:
                continue
            following_rows = candidates[position + 1 : position + 4]
            if not following_rows:
                return position
            if any(
                len(
                    populated_columns
                    & {
                        index
                        for index, value in enumerate(next_row)
                        if value not in (None, "")
                    }
                )
                >= max(1, len(populated_columns) // 2)
                for _, next_row in following_rows
            ):
                return position
        return 0

